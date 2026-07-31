"""NPC family persona Markdown files의 구조와 KB 매핑을 검증한다."""

from __future__ import annotations

import argparse
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PERSONA_DIR = ROOT / "data" / "prompts" / "npc" / "personas"

NPC_HEADING = re.compile(r"^## NPC\s+(\d+)\s+—\s+(.+)$", re.MULTILINE)
PROMPT_HEADING = "### 페르소나 프롬프트"
INCIDENT_MARKERS = (
    "### 돌발 연출 힌트",
    "**돌발 연출 힌트**",
    "돌발 연출 힌트:",
    "- 돌발 연출 힌트:",
)
F03_REQUIRED_LABELS = (
    "역할",
    "성격",
    "가치관",
    "말버릇",
    "말투",
    "담당 업무",
    "신입 대응",
    "잘했을 때",
    "실수했을 때",
    "무리한 요구",
    "권한 경계",
    "반응 원칙",
)


def find_kb_workbook() -> Path:
    candidates = [
        *ROOT.glob("*.xlsx"),
        *(ROOT / "data" / "research").glob("*.xlsx"),
        *(ROOT / "docs").glob("*.xlsx"),
    ]
    for path in candidates:
        if path.name.startswith("~$"):
            continue
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            sheet_names = workbook.sheetnames
            workbook.close()
        except Exception:
            continue
        if "01_직무마스터" in sheet_names and "02_RAG프로세스" in sheet_names:
            return path
    raise FileNotFoundError("01_직무마스터/02_RAG프로세스 시트가 있는 KB workbook을 찾지 못했습니다.")


def load_family_map(path: Path) -> dict[str, dict]:
    workbook = load_workbook(path, read_only=False, data_only=True)
    sheet = workbook["01_직무마스터"]
    families: dict[str, dict] = defaultdict(
        lambda: {"jobs": [], "titles": [], "modules": set(), "npc_roles": set()}
    )
    for row in range(2, sheet.max_row + 1):
        family_id = str(sheet.cell(row, 2).value)
        families[family_id]["jobs"].append(str(sheet.cell(row, 1).value))
        families[family_id]["titles"].append(str(sheet.cell(row, 4).value))
        families[family_id]["modules"].add(str(sheet.cell(row, 5).value))
        families[family_id]["npc_roles"].add(str(sheet.cell(row, 11).value))
    workbook.close()
    return dict(families)


def extract_prompt_blocks(text: str) -> list[str]:
    matches = list(NPC_HEADING.finditer(text))
    blocks = []
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        block = text[match.start() : end]
        safety = re.search(r"^## 공통 안전", block, re.MULTILINE)
        if safety:
            block = block[: safety.start()]
        prompt_start = block.find(PROMPT_HEADING)
        if prompt_start < 0:
            blocks.append("")
            continue
        prompt = block[prompt_start + len(PROMPT_HEADING) :].strip()
        cut_positions = [
            position
            for marker in INCIDENT_MARKERS
            if (position := prompt.find(marker)) >= 0
        ]
        if cut_positions:
            prompt = prompt[: min(cut_positions)].strip()
        blocks.append(prompt)
    return blocks


def validate_file(path: Path, family: dict) -> list[str]:
    text = path.read_text(encoding="utf-8")
    family_id = path.name[:3]
    errors: list[str] = []
    headings = NPC_HEADING.findall(text)
    prompts = extract_prompt_blocks(text)

    if not 3 <= len(headings) <= 4:
        errors.append(f"NPC 수가 3~4명이 아님: {len(headings)}")
    expected_numbers = [str(index) for index in range(1, len(headings) + 1)]
    actual_numbers = [number for number, _ in headings]
    if actual_numbers != expected_numbers:
        errors.append(f"NPC 번호가 연속적이지 않음: {actual_numbers}")
    if len(prompts) != len(headings):
        errors.append(f"프롬프트 수 불일치: NPC {len(headings)}, prompt {len(prompts)}")

    code_line = next((line for line in text.splitlines() if "codes" in line), "")
    actual_codes = re.findall(r"J\d{3}", code_line)
    if actual_codes != family["jobs"]:
        errors.append(f"KB 직무 코드 불일치: {actual_codes} != {family['jobs']}")

    if "{{ persona_prompt }}" not in text:
        errors.append("런타임 persona_prompt 슬롯 사용법 누락")
    if not re.search(r"^## 공통 안전", text, re.MULTILINE):
        errors.append("공통 안전 규칙/가드레일 누락")

    for index, prompt in enumerate(prompts, 1):
        if len(prompt) < 300:
            errors.append(f"NPC {index} 프롬프트가 300자 미만: {len(prompt)}자")
        if "당신" not in prompt:
            errors.append(f"NPC {index} 2인칭 역할 선언 누락")
        field_positions: list[int] = []
        for label in F03_REQUIRED_LABELS:
            match = re.search(rf"^- \*\*{re.escape(label)}\*\*", prompt, re.MULTILINE)
            if not match:
                errors.append(f"NPC {index} F03 필수 필드 누락: {label}")
            else:
                field_positions.append(match.start())
        if field_positions != sorted(field_positions):
            errors.append(f"NPC {index} F03 필수 필드 순서 불일치")
        has_quoted_example = bool(re.search(r'["“‘][^"”’]{4,}["”’]', prompt))
        if not has_quoted_example and not any(
            token in prompt for token in ("말투", "어투", "예:", "예시")
        ):
            errors.append(f"NPC {index} 말투 또는 발화 예시 누락")
        if not any(
            token in prompt
            for token in (
                "권한",
                "결정할 수 없",
                "할 수 없",
                "이관",
                "승인",
                "담당자",
                "책임자",
                "결정권",
                "내부 절차",
            )
        ):
            errors.append(f"NPC {index} 권한 경계 누락")
        positive = any(
            token in prompt
            for token in (
                "잘",
                "칭찬",
                "인정",
                "만족",
                "신뢰",
                "협조",
                "수긍",
                "누그러",
                "안심",
                "고마",
                "존중",
                "적극",
                "통과",
                "알려",
                "따르",
                "도와",
            )
        )
        negative = any(
            token in prompt
            for token in (
                "실수",
                "잘못",
                "누락",
                "틀리",
                "놓치",
                "위반",
                "불신",
                "항의",
                "불만",
                "거절",
                "재촉",
                "위험",
                "문제",
                "추측",
                "빠뜨",
                "모호",
                "빠진",
                "무리한",
                "보류",
                "반려",
                "정정",
                "재확인",
                "대충",
                "되묻",
                "틀린",
                "다르",
                "다를",
                "재청소",
                "얼버무리",
                "짜증",
            )
        )
        if not positive or not negative:
            errors.append(f"NPC {index} 성공/실수 반응 규칙 불완전")
        if "**NPC 역할 경계:**" not in prompt:
            errors.append(f"NPC {index} 지시 중심 역할 경계 누락")
        if "사용자의 답을 채점하거나 오류 원인을 설명" not in prompt:
            errors.append(f"NPC {index} 채점·교정 금지 규칙 누락")
        honorific_contract = (
            "자연스러운 직장 존댓말" in prompt
            and "모욕·반말" in prompt
            and "이름·직급 뒤에 ‘님’을 붙여" in prompt
            and "낮춰 부르는" in prompt
        )
        casual_contract = (
            "편한 반말" in prompt
            and "모욕·비하·인신공격" in prompt
            and "지나친 하대" in prompt
        )
        if not honorific_contract and not casual_contract:
            errors.append(f"NPC {index} 존댓말/현장 반말 상호존중 계약 누락")
        if "유자격자 또는 현장 관리자에게 보고·이관" not in prompt:
            errors.append(f"NPC {index} 위험 작업 이관 규칙 누락")

    incident_count = sum(text.count(marker) for marker in INCIDENT_MARKERS)
    if incident_count < len(headings):
        errors.append(f"돌발 연출 힌트 부족: {incident_count}/{len(headings)}")

    safety_requirements = {
        "F02": ("세무", "승인"),
        "F10": ("금융", "개인정보"),
        "F12": ("진단", "처방", "조제"),
        "F13": ("의학", "중단"),
        "F19": ("안전",),
        "F22": ("활선", "유자격자"),
        "F23": ("안전", "신고"),
        "F24": ("교통법규", "중량물"),
    }
    for token in safety_requirements.get(family_id, ()):
        if token not in text:
            errors.append(f"강화 안전 키워드 누락: {token}")

    return errors


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--kb", type=Path, default=None, help="KB v5 workbook 경로")
    args = parser.parse_args()

    kb_path = args.kb or find_kb_workbook()
    families = load_family_map(kb_path)
    paths = sorted(PERSONA_DIR.glob("F*.md"))
    expected_ids = [f"F{index:02d}" for index in range(1, 25)]
    actual_ids = [path.name[:3] for path in paths]

    all_errors: dict[str, list[str]] = {}
    if actual_ids != expected_ids:
        all_errors["FILE_SET"] = [f"family 파일 집합 불일치: {actual_ids}"]

    total_npcs = 0
    total_jobs = 0
    for path in paths:
        family_id = path.name[:3]
        if family_id not in families:
            all_errors[path.name] = ["KB에 없는 family_id"]
            continue
        text = path.read_text(encoding="utf-8")
        total_npcs += len(NPC_HEADING.findall(text))
        total_jobs += len(families[family_id]["jobs"])
        errors = validate_file(path, families[family_id])
        if errors:
            all_errors[path.name] = errors

    print(f"KB: {kb_path.name}")
    print(f"families={len(paths)} jobs={total_jobs} npc_personas={total_npcs}")
    if all_errors:
        for filename, errors in all_errors.items():
            print(f"FAIL {filename}")
            for error in errors:
                print(f"  - {error}")
        return 1

    print("PASS: 24개 family의 NPC persona 구조와 KB 직무 코드 매핑이 유효합니다.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
