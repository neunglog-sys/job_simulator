"""두 조사 엑셀과 보강 조사 YAML을 103개 직무별 NPC 오버라이드로 변환한다."""

from __future__ import annotations

import argparse
import re
from collections import defaultdict
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
RESEARCH_DIR = ROOT / "data" / "research"
OUTPUT_DIR = ROOT / "data" / "prompts" / "npc" / "jobs"
RESEARCHED_PATH = RESEARCH_DIR / "npc-job-overrides-researched.yaml"

SITUATION_TYPES = ["정상업무", "자료·정보 누락", "우선순위 충돌", "오류·안전위험", "보고·인계"]
NPC_FORBIDDEN = [
    "플레이어를 채점하거나 점수·통과 여부를 말하지 않는다.",
    "오류의 정답·모범답안·단계별 해설을 제공하지 않는다.",
    "업무상 필요한 지시·사실·자료·권한 범위만 전달한다.",
]


def _nonempty_rows(ws, start: int) -> list[tuple[Any, ...]]:
    return [
        row
        for row in ws.iter_rows(min_row=start, values_only=True)
        if any(value is not None and str(value).strip() for value in row)
    ]


def _split(value: Any) -> list[str]:
    if value is None:
        return []
    return [part.strip() for part in re.split(r"[,;/]", str(value)) if part.strip()]


def _safe_name(value: str) -> str:
    return re.sub(r'[\\/:*?"<>|]+', "_", value).strip().replace(" ", "_")


def _find_workbooks() -> tuple[Path, Path]:
    books = list(RESEARCH_DIR.glob("*.xlsx"))
    kb = next((path for path in books if "RAG_KB" in path.name), None)
    team = next((path for path in books if "RAG_KB" not in path.name), None)
    if kb is None or team is None:
        raise FileNotFoundError("data/research의 KB·통합조사 xlsx를 찾을 수 없습니다.")
    return kb, team


def _source_meta(filename: str, sheet: str, rows: str) -> dict[str, str]:
    return {"file": filename, "sheet": sheet, "rows": rows}


def _linked_variants(
    job_id: str,
    links: list[tuple[Any, ...]],
    stages_by_category: dict[str, list[tuple[Any, ...]]],
    missions_by_category: dict[str, list[tuple[Any, ...]]],
) -> list[dict[str, Any]]:
    variants: list[dict[str, Any]] = []
    for link in links:
        linked_ids = set(re.findall(r"J\d{3}", str(link[9] or "")))
        if job_id not in linked_ids:
            continue
        category = str(link[3]).strip()
        stage_rows = sorted(stages_by_category.get(category, []), key=lambda row: str(row[3]))
        mission_rows = missions_by_category.get(category, [])
        variants.append(
            {
                "category": category,
                "module": link[2],
                "rag_status": link[6],
                "workflow": link[11],
                "npc_visible": {
                    "npc_roles": _split(link[12]),
                    "input_materials": _split(link[13]),
                    "outputs": _split(link[14]),
                    "sudden_events": _split(link[16]),
                },
                "stages": [
                    {
                        "stage_id": row[3],
                        "name": row[4],
                        "goal": row[5],
                        "npc_visible": {
                            "materials": _split(row[6]),
                            "npc_roles": _split(row[7]),
                            "outputs": _split(row[8]),
                            "mission_candidate": row[9],
                        },
                        "coach_private": {
                            "success_criteria": _split(row[10]),
                            "failure_patterns": _split(row[11]),
                        },
                    }
                    for row in stage_rows
                ],
                "missions": [
                    {
                        "mission_id": row[1],
                        "situation_type": row[4],
                        "difficulty": row[5],
                        "npc_visible": {
                            "npc": row[6],
                            "request_line": row[7],
                            "materials": _split(row[8]),
                            "user_mission": row[9],
                            "outputs": _split(row[11]),
                            "sudden_event": row[14],
                        },
                        "coach_private": {
                            "required_actions": _split_numbered(row[10]),
                            "success_criteria": _split(row[12]),
                            "failure_patterns": _split(row[13]),
                        },
                        "scorer_private": {"rag_link": row[15]},
                    }
                    for row in mission_rows
                ],
            }
        )
    return variants


def _split_numbered(value: Any) -> list[str]:
    if value is None:
        return []
    parts = re.split(r"\s*[①②③④⑤⑥⑦⑧⑨⑩]\s*", str(value))
    return [part.strip() for part in parts if part.strip()]


def _researched_variant(job: dict[str, Any], constraints: list[str]) -> dict[str, Any]:
    roles = job["npc_roles"]
    tasks = job["stage_tasks"]
    stages = []
    for index, task in enumerate(tasks, 1):
        stages.append(
            {
                "stage_id": f"S{index}",
                "name": task,
                "goal": task,
                "npc_visible": {
                    "materials": job["work_targets"],
                    "npc_roles": [role["role"] for role in roles],
                    "outputs": job["outputs"],
                    "mission_candidate": task,
                },
                "coach_private": {
                    "success_criteria": job["coach_criteria"],
                    "failure_patterns": job["failure_patterns"],
                },
            }
        )

    missions = []
    for index, situation_type in enumerate(SITUATION_TYPES):
        role = roles[index % len(roles)]
        missions.append(
            {
                "mission_id": None,
                "situation_type": situation_type,
                "difficulty": "보통" if index in (0, 4) else "어려움",
                "npc_visible": {
                    "npc": role["role"],
                    "request_line": f"{tasks[index]} 업무를 진행해 주세요. {role['instruction_scope']} 기준을 확인하세요.",
                    "materials": job["work_targets"],
                    "user_mission": tasks[index],
                    "outputs": job["outputs"],
                    "sudden_event": job["failure_patterns"][index % len(job["failure_patterns"])],
                },
                "coach_private": {
                    "required_actions": [tasks[index]],
                    "success_criteria": job["coach_criteria"],
                    "failure_patterns": job["failure_patterns"],
                },
                "scorer_private": {"rag_link": "researched_override"},
            }
        )

    return {
        "category": job["focus"],
        "module": None,
        "rag_status": "researched",
        "workflow": " → ".join(tasks),
        "npc_visible": {
            "npc_roles": roles,
            "input_materials": job["work_targets"],
            "outputs": job["outputs"],
            "sudden_events": job["failure_patterns"],
            "constraints": constraints,
        },
        "stages": stages,
        "missions": missions,
    }


def generate(output_dir: Path = OUTPUT_DIR, clean: bool = True) -> dict[str, int]:
    kb_path, team_path = _find_workbooks()
    kb = load_workbook(kb_path, read_only=False, data_only=True)
    team = load_workbook(team_path, read_only=False, data_only=True)
    researched = yaml.safe_load(RESEARCHED_PATH.read_text(encoding="utf-8"))

    master_rows = _nonempty_rows(kb.worksheets[1], 2)
    rag_rows = _nonempty_rows(kb.worksheets[2], 2)
    links = _nonempty_rows(team.worksheets[1], 4)
    team_stages = _nonempty_rows(team.worksheets[2], 4)
    team_missions = _nonempty_rows(team.worksheets[5], 4)

    rag_by_job: dict[str, list[tuple[Any, ...]]] = defaultdict(list)
    for row in rag_rows:
        rag_by_job[str(row[1])].append(row)
    stages_by_category: dict[str, list[tuple[Any, ...]]] = defaultdict(list)
    for row in team_stages:
        stages_by_category[str(row[1])].append(row)
    missions_by_category: dict[str, list[tuple[Any, ...]]] = defaultdict(list)
    for row in team_missions:
        missions_by_category[str(row[3])].append(row)

    if clean and output_dir.exists():
        for path in output_dir.glob("J*.yaml"):
            path.unlink()
    output_dir.mkdir(parents=True, exist_ok=True)

    researched_count = 0
    linked_count = 0
    for master in master_rows:
        job_id, family_id, family_name, job_name, module = map(str, master[:5])
        kb_stages = sorted(rag_by_job[job_id], key=lambda row: str(row[6]))
        variants = _linked_variants(job_id, links, stages_by_category, missions_by_category)
        research_job = researched["jobs"].get(job_id)
        research_sources: list[dict[str, str]] = []
        if not variants:
            if research_job is None:
                raise ValueError(f"{job_id}: 통합조사 연결도, 보강 조사도 없습니다.")
            variants = [
                _researched_variant(
                    research_job,
                    researched["family_constraints"].get(family_id, []),
                )
            ]
            for source_id in research_job["source_ids"]:
                research_sources.append({"source_id": source_id, **researched["sources"][source_id]})
            researched_count += 1
            review_status = "researched_needs_human_review"
        else:
            linked_count += 1
            review_status = "integrated_research"

        document = {
            "schema_version": "npc.job-override.v1",
            "job": {
                "job_id": job_id,
                "job_name": job_name,
                "family_id": family_id,
                "family_name": family_name,
                "module": module,
            },
            "review": {
                "status": review_status,
                "generated": True,
                "human_review_required": review_status != "integrated_research",
            },
            "prompt_contract": {
                "npc_visible": {
                    "work_targets": _split(master[9]),
                    "npc_roles": _split(master[10]),
                    "outputs": _split(master[11]),
                    "forbidden": NPC_FORBIDDEN,
                },
                "coach_private": {
                    "success_criteria": _split(master[16]),
                    "failure_patterns": _split(master[17]),
                },
            },
            "kb_stages": [
                {
                    "chunk_id": row[0],
                    "stage_id": row[6],
                    "name": row[7],
                    "goal": row[8],
                    "rag_chunk_text": row[17],
                }
                for row in kb_stages
            ],
            "content_variants": variants,
            "provenance": {
                "excel": [
                    _source_meta(kb_path.name, "01_직무마스터", f"job_id={job_id}"),
                    _source_meta(kb_path.name, "02_RAG프로세스", f"job_id={job_id}"),
                    _source_meta(team_path.name, "01_연결맵/02_단계별프로세스/05_상황별미션", f"job_id={job_id}"),
                ],
                "research_sources": research_sources,
            },
        }
        out = output_dir / f"{job_id}_{_safe_name(job_name)}.yaml"
        out.write_text(
            yaml.safe_dump(document, allow_unicode=True, sort_keys=False, width=120),
            encoding="utf-8",
        )

    return {"total": len(master_rows), "linked": linked_count, "researched": researched_count}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    parser.add_argument("--no-clean", action="store_true")
    args = parser.parse_args()
    counts = generate(args.output_dir, clean=not args.no_clean)
    print(f"generated={counts['total']} linked={counts['linked']} researched={counts['researched']}")


if __name__ == "__main__":
    main()
